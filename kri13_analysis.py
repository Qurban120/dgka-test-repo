import os
from datetime import datetime

def analyze_kri13_from_excel(file_path):
    try:
        from openpyxl import load_workbook
        
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Find columns
        headers = [cell.value for cell in worksheet[1] if cell.value]
        issue_key_col = next((i for i, h in enumerate(headers) if 'Issue Key' in str(h)), None)
        duration_col = next((i for i, h in enumerate(headers) if 'Incident duration' in str(h)), None)
        
        if issue_key_col is None or duration_col is None:
            return []
        
        RTO_THRESHOLD = 7200  # 2 hours = 7200 seconds
        system_within_rto_counts = {}
        current_critical_system = None
        
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(issue_key_col, duration_col):
                continue
                
            issue_key = row[issue_key_col] if issue_key_col < len(row) else ""
            duration = row[duration_col] if duration_col < len(row) else ""
            
            # Check if critical system
            if issue_key and not str(issue_key).startswith('IMP-'):
                if '(' in str(issue_key) and 'ITAM-' in str(issue_key):
                    current_critical_system = str(issue_key).split('(')[0].strip()
            elif str(issue_key).startswith('IMP-') and current_critical_system:
                try:
                    duration_seconds = int(duration) if duration and str(duration).isdigit() else 0
                    
                    # Count incidents resolved WITHIN RTO (≤ 7200 seconds)
                    if 0 < duration_seconds <= RTO_THRESHOLD:
                        if current_critical_system not in system_within_rto_counts:
                            system_within_rto_counts[current_critical_system] = 0
                        system_within_rto_counts[current_critical_system] += 1
                except:
                    continue
        
        # Group Birbank systems together
        birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
        grouped_counts = {}
        
        for system, count in system_within_rto_counts.items():
            if system in birbank_systems:
                # Group all Birbank systems under "Birbank"
                if "Birbank" not in grouped_counts:
                    grouped_counts["Birbank"] = 0
                grouped_counts["Birbank"] += count
            else:
                # Keep other systems as they are
                grouped_counts[system] = count
        
        # Create results
        results = []
        for system, count in sorted(grouped_counts.items()):
            if count > 0:  # Only include systems with incidents within RTO
                results.append({
                    "System": system,
                    "Within_RTO_Count": count,
                    "Status": "EXCEEDED" if count > 2 else "WITHIN LIMIT"
                })
        
        return results
    except:
        return []

def generate_html_report(results):
    # Count systems exceeding threshold
    exceeded_count = len([r for r in results if r['Within_RTO_Count'] > 2])
    
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <title>KRI13 - RTO Within Limit Incidents Analysis</title>
    <style>
        table {{ border-collapse: collapse; width: 100%; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
        th {{ background-color: #f2f2f2; font-weight: bold; }}
        .exceeded {{ background-color: #ffcccc; color: #cc0000; font-weight: bold; }}
    </style>
</head>
<body>
    <h1>KRI13 - RTO Within Limit Incidents Analysis</h1>
    
    <p><strong>KRI13 Definition:</strong> Number of incidents in critical systems resolved within RTO</p>
    <p><strong>RTO Threshold:</strong> 2 hours (7200 seconds)</p>
    <p><strong>Maximum Allowed:</strong> 2 incidents per system (if more, threshold is violated)</p>
    
    <p style="color: red; font-weight: bold; font-size: 18px;">
        Systems exceeding threshold: {exceeded_count}
    </p>
    
    <table>
        <tr>
            <th>System</th>
            <th>Incidents Within RTO</th>
            <th>Status</th>
        </tr>"""
    
    if results:
        for result in results:
            row_class = ' class="exceeded"' if result['Within_RTO_Count'] > 2 else ''
            html_content += f"""
        <tr{row_class}>
            <td>{result['System']}</td>
            <td>{result['Within_RTO_Count']}</td>
            <td>{result['Status']}</td>
        </tr>"""
    else:
        html_content += f"""
        <tr>
            <td colspan="3" style="text-align: center; color: green; font-weight: bold;">
                ✅ No incidents resolved within RTO - All systems optimal!
            </td>
        </tr>"""
    
    html_content += """
    </table>
</body>
</html>"""
    
    return html_content

def main():
    excel_file_path = r"C:\Users\XaniyevQX\Desktop\AUTOMATION_NEW\Incident_Report_for_GRC__KRI_ (1).xlsx"
    
    if not os.path.exists(excel_file_path):
        print("Excel file not found!")
        return
    
    results = analyze_kri13_from_excel(excel_file_path)
    
    html_content = generate_html_report(results)
    
    with open('KRI13_RTO_Within_Report.html', 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print("Successfully reported!")

if __name__ == "__main__":
    main()