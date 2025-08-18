import os
import io
from typing import Optional, Dict, Any, List

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel

from db import (
    init_db,
    seed_from_json,
    fetch_all_data,
    get_kri,
    add_system as db_add_system,
    upsert_data_point,
    update_data_point as db_update_data_point,
    delete_data_point as db_delete_data_point,
    delete_system as db_delete_system,
    VALUE_FIELD_MAP,
    APP_BASED_KRIS,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INDEX_FILE = os.path.join(BASE_DIR, 'index.html')


# Pydantic models
class AddSystemRequest(BaseModel):
    kriId: str
    systemName: str


class AddDataPointRequest(BaseModel):
    kriId: str
    period: str
    value: float
    valueField: Optional[str] = None
    systemName: Optional[str] = None


class UpdateDataPointRequest(BaseModel):
    kriId: str
    period: str
    systemName: Optional[str] = None
    value: Optional[float] = None
    valueField: Optional[str] = None


class DeleteDataPointRequest(BaseModel):
    kriId: str
    period: str
    systemName: Optional[str] = None


class DeleteSystemRequest(BaseModel):
    kriId: str
    systemName: str


class ExportRequest(BaseModel):
    kriId: str
    format: str


app = FastAPI()

# CORS (relaxed for simplicity; tighten in prod)
app.add_middleware(
    CORSMiddleware,
    allow_origins=['*'],
    allow_credentials=True,
    allow_methods=['*'],
    allow_headers=['*'],
)


@app.on_event('startup')
async def startup_event():
    # Ensure DB exists and seed from JSON if empty
    init_db(drop=False)
    # Optionally seed when DB is empty (first run)
    data = fetch_all_data()
    if all((not v.get('data')) and (not v.get('applications')) for v in data.values()):
        seed_from_json()


@app.get('/', response_class=HTMLResponse)
async def get_index():
    if not os.path.exists(INDEX_FILE):
        raise HTTPException(status_code=404, detail='index.html not found')
    with open(INDEX_FILE, 'r', encoding='utf-8') as f:
        return HTMLResponse(content=f.read(), status_code=200)


@app.get('/api/kri-data')
async def api_get_kri_data():
    return fetch_all_data()


@app.get('/api/kri/{kri_id}')
async def api_get_single_kri(kri_id: str):
    if kri_id not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    return get_kri(kri_id)


@app.post('/api/add-system')
async def api_add_system(req: AddSystemRequest):
    kri = req.kriId
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    if kri not in APP_BASED_KRIS:
        raise HTTPException(status_code=400, detail='KRI is not applications-based')
    name = req.systemName.strip()
    if not name:
        raise HTTPException(status_code=400, detail='System name is required')
    db_add_system(kri, name)
    return {'success': True, 'kriId': kri, 'kri': get_kri(kri)}


@app.post('/api/add-data-point')
async def api_add_data_point(req: AddDataPointRequest):
    kri = req.kriId
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')

    field = req.valueField or VALUE_FIELD_MAP[kri]
    if field != VALUE_FIELD_MAP[kri]:
        raise HTTPException(status_code=400, detail=f'Invalid valueField for {kri}. Expected {VALUE_FIELD_MAP[kri]}')

    period = req.period.strip()
    if not period:
        raise HTTPException(status_code=400, detail='Period is required')

    value = float(req.value)
    if value < 0:
        raise HTTPException(status_code=400, detail='Value must be non-negative')

    if kri in APP_BASED_KRIS and not req.systemName:
        raise HTTPException(status_code=400, detail='systemName is required for applications-based KRI')

    upsert_data_point(kri, period, value, system_name=req.systemName)
    return {'success': True, 'kriId': kri, 'kri': get_kri(kri)}


@app.put('/api/update-data-point')
async def api_update_data_point(req: UpdateDataPointRequest):
    kri = req.kriId
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    if req.valueField and req.valueField != VALUE_FIELD_MAP[kri]:
        raise HTTPException(status_code=400, detail=f'Invalid valueField for {kri}. Expected {VALUE_FIELD_MAP[kri]}')

    changed = db_update_data_point(
        kri_id=kri,
        period=req.period,
        value=req.value,
        system_name=req.systemName,
    )
    if changed == 0:
        raise HTTPException(status_code=404, detail='Data point not found or no changes provided')
    return {'success': True, 'kriId': kri, 'kri': get_kri(kri)}


@app.delete('/api/delete-data-point')
async def api_delete_data_point(req: DeleteDataPointRequest):
    kri = req.kriId
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    deleted = db_delete_data_point(kri, req.period, req.systemName)
    if deleted == 0:
        raise HTTPException(status_code=404, detail='Data point not found')
    return {'success': True, 'kriId': kri, 'kri': get_kri(kri)}


@app.delete('/api/delete-system')
async def api_delete_system(req: DeleteSystemRequest):
    kri = req.kriId
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    if kri not in APP_BASED_KRIS:
        raise HTTPException(status_code=400, detail='KRI is not applications-based')
    pts, sys = db_delete_system(kri, req.systemName)
    if sys == 0:
        raise HTTPException(status_code=404, detail='System not found')
    return {'success': True, 'deletedPoints': pts, 'deletedSystems': sys, 'kriId': kri, 'kri': get_kri(kri)}


def _kri_meta(kri_id: str) -> Dict[str, str]:
    definitions = {
        'kri2': 'Percentage of critical initiatives/projects engage with information security at the requirements phase',
        'kri3': 'Percentage of environments evaluated to determine whether or not they are critical',
        'kri4': 'Number of realized risk which is coming from threats that identified as weak threat with low probability',
        'kri5': 'Percentage of outstanding Critical risk mitigation actions vs. agreed plan',
        'kri6': 'Percentage of outstanding High and moderate risk mitigation actions vs. agreed plan',
        'kri8': 'Mean Time Between Critical System Failures (Days)',
        'kri9': 'Unauthorized access to strictly confidential data in any critical systems',
        'kri10': 'Number of incidents affecting critical applications',
        'kri12': 'Incidents Exceeding RTO',
        'kri13': 'Number of incidents resulting in downtime within RTO',
        'kri15': '% of Critical Systems without up to date patches',
        'kri18': 'Percentage of critical systems not included in BCP/DR plans',
        'kri19': 'Average time to recover from an incident',
        'kri20': 'Number of incidents resulting in loss of customer information',
        'kri21': 'Number of privacy incidents resulting in legal/regulatory proceedings',
        'kri22': 'Percentage of incidents detected by controls',
        'kri27': 'Critical risks with no owner and action plan',
        'kri28': 'High/Moderate risks with no owner and action plan',
    }
    thresholds = {
        'kri2': '95% (higher is better)',
        'kri3': '80% (higher is better)',
        'kri4': '1 realized risk (lower is better)',
        'kri5': '0% for Critical Risks (lower is better)',
        'kri6': '10% (lower is better)',
        'kri8': '120 days (higher is better)',
        'kri9': '0 (lower is better)',
        'kri10': '2 incidents per month (lower is better)',
        'kri12': '0 (lower is better)',
        'kri13': '2 (lower is better)',
        'kri15': '0% (lower is better)',
        'kri18': '10% (lower is better)',
        'kri19': '< RTO time (BCP) (lower is better)',
        'kri20': '0 incidents per month (lower is better)',
        'kri21': '0 incidents per month (lower is better)',
        'kri22': '70% (higher is better)',
        'kri27': '0% for Critical Risks (lower is better)',
        'kri28': '10% for High/Moderate (lower is better)',
    }
    return {
        'name': kri_id.upper(),
        'definition': definitions.get(kri_id, ''),
        'threshold': thresholds.get(kri_id, ''),
    }


def _generate_html_report(kri_id: str, data: Dict[str, Any]) -> str:
    meta = _kri_meta(kri_id)
    def render_points() -> str:
        if 'data' in data:
            rows = ''.join(f"<tr><td>{p['period']}</td><td>{list(p.values())[1]}</td></tr>" for p in data['data'])
            return f"<table border='1' cellspacing='0' cellpadding='6'><thead><tr><th>Period</th><th>Value</th></tr></thead><tbody>{rows}</tbody></table>"
        apps = data.get('applications', {})
        parts: List[str] = []
        for app, items in apps.items():
            rows = ''.join(f"<tr><td>{i['period']}</td><td>{list(i.values())[1]}</td></tr>" for i in items)
            parts.append(f"<h3>{app}</h3><table border='1' cellspacing='0' cellpadding='6'><thead><tr><th>Period</th><th>Value</th></tr></thead><tbody>{rows}</tbody></table>")
        return ''.join(parts) or '<p>No data</p>'

    return f"""
<!DOCTYPE html>
<html>
<head>
  <meta charset='utf-8'>
  <title>{meta['name']} Report</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 24px; color:#222; }}
    h1 {{ color:#333; }}
    .meta {{ margin-bottom:16px; }}
    table {{ border-collapse: collapse; margin-bottom: 16px; width: 100%; }}
    th, td {{ border: 1px solid #ccc; padding: 8px; text-align: left; }}
    th {{ background:#f2f3f7; }}
  </style>
}</head>
<body>
  <h1>{meta['name']} Report</h1>
  <div class='meta'>
    <p><strong>Definition:</strong> {meta['definition']}</p>
    <p><strong>Threshold:</strong> {meta['threshold']}</p>
  </div>
  <h2>Data</h2>
  {render_points()}
</body>
</html>
"""


def _generate_txt_report(kri_id: str, data: Dict[str, Any]) -> str:
    meta = _kri_meta(kri_id)
    lines = [
        f"{meta['name']} Report",
        f"Definition: {meta['definition']}",
        f"Threshold: {meta['threshold']}",
        '',
        'Data:',
    ]
    if 'data' in data:
        for p in data['data']:
            vals = list(p.values())
            lines.append(f"- {p['period']}: {vals[1]}")
    else:
        apps = data.get('applications', {})
        for app, items in apps.items():
            lines.append(f"{app}:")
            for i in items:
                vals = list(i.values())
                lines.append(f"  - {i['period']}: {vals[1]}")
    return '\n'.join(lines)


def _generate_docx_report(kri_id: str, data: Dict[str, Any]) -> bytes:
    try:
        from docx import Document  # python-docx
    except Exception as e:
        raise HTTPException(status_code=500, detail=f'DOCX generation not available: {e}')
    meta = _kri_meta(kri_id)
    doc = Document()
    doc.add_heading(f"{meta['name']} Report", level=1)
    doc.add_paragraph(f"Definition: {meta['definition']}")
    doc.add_paragraph(f"Threshold: {meta['threshold']}")
    doc.add_heading('Data', level=2)
    if 'data' in data:
        table = doc.add_table(rows=1, cols=2)
        hdr = table.rows[0].cells
        hdr[0].text = 'Period'
        hdr[1].text = 'Value'
        for p in data['data']:
            row = table.add_row().cells
            row[0].text = str(p['period'])
            vals = list(p.values())
            row[1].text = str(vals[1])
    else:
        apps = data.get('applications', {})
        if not apps:
            doc.add_paragraph('No data')
        for app, items in apps.items():
            doc.add_heading(app, level=3)
            table = doc.add_table(rows=1, cols=2)
            hdr = table.rows[0].cells
            hdr[0].text = 'Period'
            hdr[1].text = 'Value'
            for i in items:
                row = table.add_row().cells
                row[0].text = str(i['period'])
                vals = list(i.values())
                row[1].text = str(vals[1])
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()


def _generate_pdf_report(kri_id: str, data: Dict[str, Any]) -> bytes:
    # Simple approach: convert generated HTML to PDF using WeasyPrint if available; fallback to plain text PDF via reportlab
    html = _generate_html_report(kri_id, data)
    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html).write_pdf()
        return pdf_bytes
    except Exception:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
            from reportlab.lib.units import cm
        except Exception as e:
            raise HTTPException(status_code=500, detail=f'PDF generation not available: {e}')
        txt = _generate_txt_report(kri_id, data)
        bio = io.BytesIO()
        c = canvas.Canvas(bio, pagesize=A4)
        width, height = A4
        x, y = 2*cm, height - 2*cm
        for line in txt.split('\n'):
            c.drawString(x, y, line[:110])
            y -= 14
            if y < 2*cm:
                c.showPage()
                y = height - 2*cm
        c.save()
        bio.seek(0)
        return bio.read()


@app.post('/api/export-report')
async def api_export_report(req: ExportRequest):
    kri = req.kriId
    fmt = (req.format or 'html').lower()
    if kri not in VALUE_FIELD_MAP:
        raise HTTPException(status_code=400, detail='Unknown KRI id')
    data = get_kri(kri)
    filename = f"{kri}-report.{fmt}"
    if fmt == 'html':
        html = _generate_html_report(kri, data)
        return HTMLResponse(content=html, headers={
            'Content-Disposition': f'attachment; filename="{filename}"'
        })
    if fmt == 'txt':
        txt = _generate_txt_report(kri, data).encode('utf-8')
        return StreamingResponse(io.BytesIO(txt), media_type='text/plain', headers={
            'Content-Disposition': f'attachment; filename="{filename}"'
        })
    if fmt == 'docx':
        content = _generate_docx_report(kri, data)
        return StreamingResponse(io.BytesIO(content), media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', headers={
            'Content-Disposition': f'attachment; filename="{filename}"'
        })
    if fmt == 'pdf':
        content = _generate_pdf_report(kri, data)
        return StreamingResponse(io.BytesIO(content), media_type='application/pdf', headers={
            'Content-Disposition': f'attachment; filename="{filename}"'
        })
    raise HTTPException(status_code=400, detail='Unsupported format')

