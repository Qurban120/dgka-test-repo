import os
from datetime import datetime

def analyze_all_kris_from_excel(file_path):
    try:
        from openpyxl import load_workbook
        
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Find columns
        headers = [cell.value for cell in worksheet[1] if cell.value]
        issue_key_col = next((i for i, h in enumerate(headers) if 'Issue Key' in str(h)), None)
        start_date_col = next((i for i, h in enumerate(headers) if 'Incident start date' in str(h)), None)
        duration_col = next((i for i, h in enumerate(headers) if 'Incident duration' in str(h)), None)
        
        if issue_key_col is None or start_date_col is None or duration_col is None:
            return {}, [], [], []
        
        RTO_THRESHOLD = 7200  # 2 hours = 7200 seconds
        
        # Data structures for different KRIs
        system_month_counts = {}  # KRI10
        rto_exceeded_incidents = []  # KRI12
        system_within_rto_counts = {}  # KRI13
        system_durations = {}  # KRI19
        
        current_critical_system = None
        
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(issue_key_col, start_date_col, duration_col):
                continue
                
            issue_key = row[issue_key_col] if issue_key_col < len(row) else ""
            start_date = row[start_date_col] if start_date_col < len(row) else ""
            duration = row[duration_col] if duration_col < len(row) else ""
            
            # Check if critical system
            if issue_key and not str(issue_key).startswith('IMP-'):
                if '(' in str(issue_key) and 'ITAM-' in str(issue_key):
                    current_critical_system = str(issue_key).split('(')[0].strip()
            elif str(issue_key).startswith('IMP-') and current_critical_system:
                
                # KRI10: Monthly incident counts
                month = None
                try:
                    if hasattr(start_date, 'month'):
                        month_num = start_date.month
                    else:
                        date_str = str(start_date).strip()
                        if date_str and date_str != "-":
                            date_part = date_str.split()[0]
                            day, month_part, year = date_part.split('/')
                            month_num = int(month_part)
                    
                    month_mapping = {4: "April", 5: "May", 6: "June"}
                    month = month_mapping.get(month_num)
                except:
                    pass
                
                if month:
                    key = (current_critical_system, month)
                    if key not in system_month_counts:
                        system_month_counts[key] = 0
                    system_month_counts[key] += 1
                
                # KRI12, KRI13, KRI19: Duration-based analysis
                try:
                    duration_seconds = int(duration) if duration and str(duration).isdigit() else 0
                    
                    if duration_seconds > 0:
                        # KRI19: Collect all durations
                        if current_critical_system not in system_durations:
                            system_durations[current_critical_system] = []
                        system_durations[current_critical_system].append(duration_seconds)
                        
                        # KRI12: RTO exceeded incidents
                        if duration_seconds > RTO_THRESHOLD:
                            rto_exceeded_incidents.append({
                                "System": current_critical_system,
                                "Incident": issue_key,
                                "RTO_Exceeded_Seconds": duration_seconds,
                                "RTO_Exceeded_Hours": round(duration_seconds / 3600, 2)
                            })
                        
                        # KRI13: Within RTO incidents
                        if duration_seconds <= RTO_THRESHOLD:
                            if current_critical_system not in system_within_rto_counts:
                                system_within_rto_counts[current_critical_system] = 0
                            system_within_rto_counts[current_critical_system] += 1
                except:
                    continue
        
        return system_month_counts, rto_exceeded_incidents, system_within_rto_counts, system_durations
        
    except:
        return {}, [], [], []

def process_kri10(system_month_counts):
    # Group Birbank systems
    birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
    grouped_counts = {}
    
    for (system, month), count in system_month_counts.items():
        if system in birbank_systems:
            key = ("Birbank", month)
        else:
            key = (system, month)
        
        if key not in grouped_counts:
            grouped_counts[key] = 0
        grouped_counts[key] += count
    
    results = []
    critical_systems = set([key[0] for key in grouped_counts.keys()])
    
    for system in sorted(critical_systems):
        for month in ["April", "May", "June"]:
            count = grouped_counts.get((system, month), 0)
            if count > 0:
                results.append({
                    "System": system,
                    "Month": month,
                    "Count": count,
                    "Status": "EXCEEDED" if count > 2 else "WITHIN LIMIT"
                })
    
    return results

def process_kri12(rto_exceeded_incidents):
    # Group Birbank systems
    birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
    grouped_incidents = []
    
    for incident in rto_exceeded_incidents:
        if incident["System"] in birbank_systems:
            incident["System"] = "Birbank"
        grouped_incidents.append(incident)
    
    return grouped_incidents

def process_kri13(system_within_rto_counts):
    # Group Birbank systems
    birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
    grouped_counts = {}
    
    for system, count in system_within_rto_counts.items():
        if system in birbank_systems:
            if "Birbank" not in grouped_counts:
                grouped_counts["Birbank"] = 0
            grouped_counts["Birbank"] += count
        else:
            grouped_counts[system] = count
    
    results = []
    for system, count in sorted(grouped_counts.items()):
        if count > 0:
            results.append({
                "System": system,
                "Count": count,
                "Status": "EXCEEDED" if count > 2 else "WITHIN LIMIT"
            })
    
    return results

def process_kri19(system_durations):
    # Group Birbank systems
    birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
    grouped_durations = {}
    
    for system, durations in system_durations.items():
        if system in birbank_systems:
            if "Birbank" not in grouped_durations:
                grouped_durations["Birbank"] = []
            grouped_durations["Birbank"].extend(durations)
        else:
            grouped_durations[system] = durations
    
    results = []
    for system, durations in sorted(grouped_durations.items()):
        if durations:
            total_duration = sum(durations)
            incident_count = len(durations)
            average_duration = total_duration / incident_count
            
            results.append({
                "System": system,
                "Incident_Count": incident_count,
                "Average_Seconds": int(average_duration),
                "Average_Hours": round(average_duration / 3600, 2),
                "Status": "EXCEEDED RTO" if average_duration > 7200 else "WITHIN RTO"
            })
    
    return results

def generate_integrated_html_report(kri10_results, kri12_results, kri13_results, kri19_results):
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <title>Integrated KRI Analysis Report - Q2 2025</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; background-color: #f5f5f5; }}
        .container {{ max-width: 1200px; margin: 0 auto; background-color: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 10px rgba(0,0,0,0.1); }}
        h1 {{ color: #2c3e50; text-align: center; margin-bottom: 30px; }}
        h2 {{ color: #34495e; border-bottom: 2px solid #3498db; padding-bottom: 10px; margin-top: 40px; }}
        table {{ border-collapse: collapse; width: 100%; margin-bottom: 30px; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
        th {{ background-color: #34495e; color: white; font-weight: bold; }}
        tr:nth-child(even) {{ background-color: #f2f2f2; }}
        .exceeded {{ background-color: #ffcccc; color: #cc0000; font-weight: bold; }}
        .summary {{ background-color: #e8f4f8; padding: 15px; margin-bottom: 20px; border-left: 4px solid #3498db; }}
        .kri-summary {{ display: flex; justify-content: space-around; margin-bottom: 30px; }}
        .kri-box {{ background-color: #ecf0f1; padding: 15px; border-radius: 5px; text-align: center; }}
        .exceeded-count {{ color: red; font-weight: bold; font-size: 18px; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>Integrated KRI Analysis Report</h1>
        <h3 style="text-align: center; color: #666;">Q2 2025 (April - June)</h3>
        
        <div class="kri-summary">
            <div class="kri-box">
                <h4>KRI10</h4>
                <div class="exceeded-count">{len([r for r in kri10_results if r['Count'] > 2])}</div>
                <p>Systems exceeding monthly threshold</p>
            </div>
            <div class="kri-box">
                <h4>KRI12</h4>
                <div class="exceeded-count">{len(kri12_results)}</div>
                <p>Incidents exceeding RTO</p>
            </div>
            <div class="kri-box">
                <h4>KRI13</h4>
                <div class="exceeded-count">{len([r for r in kri13_results if r['Count'] > 2])}</div>
                <p>Systems with too many within-RTO incidents</p>
            </div>
            <div class="kri-box">
                <h4>KRI19</h4>
                <div class="exceeded-count">{len([r for r in kri19_results if r['Average_Seconds'] > 7200])}</div>
                <p>Systems with average time > RTO</p>
            </div>
        </div>"""
    
    # KRI10 Section
    html_content += f"""
        <h2>KRI10 - Monthly Incident Counts</h2>
        <div class="summary">
            <strong>Definition:</strong> Number of incidents affecting critical systems per month<br>
            <strong>Threshold:</strong> Maximum 2 incidents per system per month
        </div>
        <table>
            <tr>
                <th>System</th>
                <th>Month</th>
                <th>Count of Incidents</th>
                <th>Status</th>
            </tr>"""
    
    for result in kri10_results:
        row_class = ' class="exceeded"' if result['Count'] > 2 else ''
        html_content += f"""
            <tr{row_class}>
                <td>{result['System']}</td>
                <td>{result['Month']}</td>
                <td>{result['Count']}</td>
                <td>{result['Status']}</td>
            </tr>"""
    
    # KRI12 Section
    html_content += f"""
        </table>
        
        <h2>KRI12 - RTO Exceeded Incidents</h2>
        <div class="summary">
            <strong>Definition:</strong> Incidents resolved longer than RTO (>2 hours)<br>
            <strong>Target:</strong> 0 incidents exceeding RTO for each system
        </div>
        <table>
            <tr>
                <th>System</th>
                <th>Incident</th>
                <th>Duration (Seconds)</th>
                <th>Duration (Hours)</th>
            </tr>"""
    
    if kri12_results:
        for result in kri12_results:
            html_content += f"""
            <tr class="exceeded">
                <td>{result['System']}</td>
                <td>{result['Incident']}</td>
                <td>{result['RTO_Exceeded_Seconds']}</td>
                <td>{result['RTO_Exceeded_Hours']}</td>
            </tr>"""
    else:
        html_content += """
            <tr>
                <td colspan="4" style="text-align: center; color: green; font-weight: bold;">
                    ✅ No incidents exceeded RTO threshold
                </td>
            </tr>"""
    
    # KRI13 Section
    html_content += f"""
        </table>
        
        <h2>KRI13 - Within RTO Incident Counts</h2>
        <div class="summary">
            <strong>Definition:</strong> Number of incidents resolved within RTO (≤2 hours)<br>
            <strong>Threshold:</strong> Maximum 2 incidents per system (if more, too many incidents occurring)
        </div>
        <table>
            <tr>
                <th>System</th>
                <th>Incidents Within RTO</th>
                <th>Status</th>
            </tr>"""
    
    for result in kri13_results:
        row_class = ' class="exceeded"' if result['Count'] > 2 else ''
        html_content += f"""
            <tr{row_class}>
                <td>{result['System']}</td>
                <td>{result['Count']}</td>
                <td>{result['Status']}</td>
            </tr>"""
    
    # KRI19 Section
    html_content += f"""
        </table>
        
        <h2>KRI19 - Average Resolution Time</h2>
        <div class="summary">
            <strong>Definition:</strong> Average incident resolution time per system<br>
            <strong>Target:</strong> Average should be less than RTO (2 hours = 7200 seconds)
        </div>
        <table>
            <tr>
                <th>System</th>
                <th>Total Incidents</th>
                <th>Average Duration (Seconds)</th>
                <th>Average Duration (Hours)</th>
                <th>Status</th>
            </tr>"""
    
    for result in kri19_results:
        row_class = ' class="exceeded"' if result['Average_Seconds'] > 7200 else ''
        html_content += f"""
            <tr{row_class}>
                <td>{result['System']}</td>
                <td>{result['Incident_Count']}</td>
                <td>{result['Average_Seconds']}</td>
                <td>{result['Average_Hours']}</td>
                <td>{result['Status']}</td>
            </tr>"""
    
    html_content += f"""
        </table>
        
        <div style="margin-top: 30px; text-align: center; color: #666; font-size: 12px;">
            Report generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
        </div>
    </div>
</body>
</html>"""
    
    return html_content

def main():
    excel_file_path = r"C:\Users\XaniyevQX\Desktop\AUTOMATION_NEW\Incident_Report_for_GRC__KRI_ (1).xlsx"
    
    if not os.path.exists(excel_file_path):
        print("Excel file not found!")
        return
    
    # Analyze all KRIs
    system_month_counts, rto_exceeded_incidents, system_within_rto_counts, system_durations = analyze_all_kris_from_excel(excel_file_path)
    
    # Process each KRI
    kri10_results = process_kri10(system_month_counts)
    kri12_results = process_kri12(rto_exceeded_incidents)
    kri13_results = process_kri13(system_within_rto_counts)
    kri19_results = process_kri19(system_durations)
    
    # Generate integrated HTML report
    html_content = generate_integrated_html_report(kri10_results, kri12_results, kri13_results, kri19_results)
    
    # Write to file
    with open('KRI_Integrated_Report.html', 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print("Successfully reported!")

if __name__ == "__main__":
    main()