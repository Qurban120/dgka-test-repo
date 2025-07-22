import os
from datetime import datetime

def extract_month_from_date(date_obj):
    """Extract month from date object or string in format DD/MM/YY HH:MM:SS"""
    if not date_obj or date_obj == "-":
        return None
    
    try:
        # Check if it's already a datetime object
        if hasattr(date_obj, 'month'):
            month_num = date_obj.month
        else:
            # Convert to string and strip whitespace
            date_str = str(date_obj).strip()
            if date_str == "" or date_str == "-":
                return None
            
            # Extract the date part (before space)
            date_part = date_str.split()[0]
            # Parse DD/MM/YY format
            day, month, year = date_part.split('/')
            month_num = int(month)
        
        # Map month numbers to names for Q2
        month_mapping = {
            4: "April",
            5: "May", 
            6: "June"
        }
        
        return month_mapping.get(month_num)
    except:
        return None

def extract_critical_system_from_issue_key(issue_key):
    """Extract critical system name from Issue Key column"""
    if not issue_key or issue_key.strip() == "":
        return None
    
    issue_key = str(issue_key).strip()
    
    # Skip IMP- entries, they are individual incidents
    if issue_key.startswith('IMP-'):
        return None
    
    # Extract system name before (ITAM-...)
    if '(' in issue_key and 'ITAM-' in issue_key:
        return issue_key.split('(')[0].strip()
    
    return issue_key

def analyze_kri10_from_excel(file_path):
    """Analyze KRI10 from Excel file directly"""
    
    try:
        # Try to import openpyxl
        try:
            from openpyxl import load_workbook
        except ImportError:
            print("❌ openpyxl library not found!")
            print("Please install it using: pip install openpyxl")
            return []
        
        # Load Excel workbook
        print(f"📖 Loading Excel file: {file_path}")
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Get headers from first row
        headers = []
        for cell in worksheet[1]:
            if cell.value:
                headers.append(cell.value)
        
        print(f"Excel columns found: {headers}")
        
        # Find column indices
        issue_key_col = None
        start_date_col = None
        
        for i, header in enumerate(headers):
            if 'Issue Key' in str(header):
                issue_key_col = i
            elif 'Incident start date' in str(header):
                start_date_col = i
        
        if issue_key_col is None or start_date_col is None:
            print("❌ Required columns not found!")
            return []
        
        system_month_counts = {}
        current_critical_system = None
        
        row_count = 0
        incident_count = 0
        
        # Process each row (starting from row 2, skipping header)
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            row_count += 1
            
            if len(row) <= max(issue_key_col, start_date_col):
                continue
                
            issue_key = row[issue_key_col] if issue_key_col < len(row) else ""
            start_date = row[start_date_col] if start_date_col < len(row) else ""
            
            # Check if this row defines a critical system
            critical_system = extract_critical_system_from_issue_key(issue_key)
            
            if critical_system:
                # This is a critical system definition row
                current_critical_system = critical_system
                print(f"Found critical system: {critical_system}")
            elif str(issue_key).startswith('IMP-') and current_critical_system:
                # This is an incident belonging to the current critical system
                incident_count += 1
                
                # Extract month from start date
                month = extract_month_from_date(start_date)
                
                # Only count incidents in Q2 months (April, May, June)
                if month:
                    key = (current_critical_system, month)
                    if key not in system_month_counts:
                        system_month_counts[key] = 0
                    system_month_counts[key] += 1
        
        print(f"\nProcessed {row_count} rows, found {incident_count} incidents")
        
        # Create results list
        results = []
        critical_systems = set([key[0] for key in system_month_counts.keys()])
        
        # For each critical system, check all Q2 months
        for system in sorted(critical_systems):
            for month in ["April", "May", "June"]:
                count = system_month_counts.get((system, month), 0)
                if count > 0:  # Only include months with incidents
                    results.append({
                        "System": system,
                        "Month": month,
                        "Count of Incidents": count
                    })
        
        return results
        
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        import traceback
        traceback.print_exc()
        return []

def generate_simple_html_report(results):
    """Generate simple HTML report for KRI10 analysis"""
    
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <title>KRI10 - Critical Systems Incident Analysis Q2 2025</title>
</head>
<body>
    <h1>KRI10 - Critical Systems Incident Analysis</h1>
    <h2>Q2 2025 (April - June)</h2>
    
    <p><strong>KRI10 Definition:</strong> Number of incidents affecting critical systems</p>
    <p><strong>Threshold:</strong> Maximum 2 incidents per critical system per month</p>
    
    <table border="1" cellpadding="5" cellspacing="0">
        <tr>
            <th>System</th>
            <th>Month</th>
            <th>Count of Incidents</th>
        </tr>"""
    
    # Add table rows
    for result in results:
        html_content += f"""
        <tr>
            <td>{result['System']}</td>
            <td>{result['Month']}</td>
            <td>{result['Count of Incidents']}</td>
        </tr>"""
    
    html_content += f"""
    </table>
    
    <h3>Summary</h3>
    <p>Total incidents analyzed: {sum(r['Count of Incidents'] for r in results)}</p>
    <p>Critical systems monitored: {len(set(r['System'] for r in results))}</p>
    <p>Report generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
    
</body>
</html>"""
    
    return html_content

def main():
    """Main function to run KRI10 analysis"""
    
    # Excel file path (as specified by user)
    excel_file_path = r"C:\Users\XaniyevQX\Desktop\AUTOMATION_NEW\Incident_Report_for_GRC__KRI_ (1).xlsx"
    
    print("KRI10 - Critical Systems Incident Analysis")
    print("=" * 50)
    print(f"Excel file path: {excel_file_path}")
    
    # Check if Excel file exists
    if not os.path.exists(excel_file_path):
        print(f"❌ Excel file not found at: {excel_file_path}")
        print("Please check the file path and try again.")
        return
    
    print("✅ Excel file found!")
    
    # Analyze the data directly from Excel
    results = analyze_kri10_from_excel(excel_file_path)
    
    if not results:
        print("❌ No data found or error occurred!")
        return
    
    # Print results to console
    print("\n" + "=" * 60)
    print("KRI10 RESULTS")
    print("=" * 60)
    print(f"{'System':<25} {'Month':<10} {'Count'}")
    print("-" * 60)
    
    for result in results:
        print(f"{result['System']:<25} {result['Month']:<10} {result['Count of Incidents']}")
    
    # Generate simple HTML report
    html_content = generate_simple_html_report(results)
    
    # Write to file
    output_file = 'KRI10_Report.html'
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print("\n" + "=" * 60)
    print("SUMMARY")
    print("=" * 60)
    print(f"✅ HTML report generated: {output_file}")
    print(f"📊 Total incidents analyzed: {sum(r['Count of Incidents'] for r in results)}")
    print(f"🏢 Critical systems monitored: {len(set(r['System'] for r in results))}")
    
    # Show systems exceeding threshold
    exceeded_systems = [r for r in results if r['Count of Incidents'] > 2]
    if exceeded_systems:
        print(f"\n⚠️  Systems exceeding KRI10 threshold (>2 incidents/month):")
        for system in exceeded_systems:
            print(f"   - {system['System']} in {system['Month']}: {system['Count of Incidents']} incidents")
    else:
        print(f"\n✅ All critical systems are within KRI10 threshold!")

if __name__ == "__main__":
    main()