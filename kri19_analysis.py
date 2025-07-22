import os
from datetime import datetime

def analyze_kri19_from_excel(file_path):
    try:
        from openpyxl import load_workbook
        
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        
        # Find columns
        headers = [cell.value for cell in worksheet[1] if cell.value]
        issue_key_col = next((i for i, h in enumerate(headers) if 'Issue Key' in str(h)), None)
        duration_col = next((i for i, h in enumerate(headers) if 'Incident duration' in str(h)), None)
        
        if issue_key_col is None or duration_col is None:
            return []
        
        RTO_THRESHOLD = 120  # 2 hours = 120 minutes
        system_durations = {}
        current_critical_system = None
        
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if len(row) <= max(issue_key_col, duration_col):
                continue
                
            issue_key = row[issue_key_col] if issue_key_col < len(row) else ""
            duration = row[duration_col] if duration_col < len(row) else ""
            
            # Check if critical system
            if issue_key and not str(issue_key).startswith('IMP-'):
                if '(' in str(issue_key) and 'ITAM-' in str(issue_key):
                    current_critical_system = str(issue_key).split('(')[0].strip()
            elif str(issue_key).startswith('IMP-') and current_critical_system:
                try:
                    duration_minutes = int(duration) if duration and str(duration).isdigit() else 0
                    
                    # Collect all incident durations for each system
                    if duration_minutes > 0:
                        if current_critical_system not in system_durations:
                            system_durations[current_critical_system] = []
                        system_durations[current_critical_system].append(duration_minutes)
                except:
                    continue
        
        # Group Birbank systems (except BirBank-Business)
        birbank_systems = ["BirBank.EDV", "BirBank.Loyalty", "BirBank.Payments", "BirBank.Transfers", "Birbank"]
        grouped_durations = {}
        
        for system, durations in system_durations.items():
            if system in birbank_systems:
                if "Birbank" not in grouped_durations:
                    grouped_durations["Birbank"] = []
                grouped_durations["Birbank"].extend(durations)
            else:
                grouped_durations[system] = durations
        
        # Calculate average resolution time for each system
        results = []
        for system, durations in sorted(grouped_durations.items()):
            if durations:
                total_duration = sum(durations)
                incident_count = len(durations)
                average_duration = total_duration / incident_count
                
                results.append({
                    "System": system,
                    "Total_Incidents": incident_count,
                    "Average_Duration_Minutes": int(average_duration),
                    "Average_Duration_Hours": round(average_duration / 60, 2),
                    "Status": "EXCEEDED RTO" if average_duration > RTO_THRESHOLD else "WITHIN RTO"
                })
        
        return results
        
    except:
        return []

def generate_kri19_html_report(results):
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <title>KRI19 - Average Resolution Time Report</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; }}
        table {{ border-collapse: collapse; width: 100%; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
        th {{ background-color: #f2f2f2; }}
        .exceeded {{ background-color: #ffcccc; }}
    </style>
</head>
<body>
    <h1>KRI19 - Average Resolution Time</h1>
    <table>
        <tr>
            <th>System</th>
            <th>Total Incidents</th>
            <th>Average Duration (Minutes)</th>
            <th>Average Duration (Hours)</th>
            <th>Status</th>
        </tr>"""
    
    for result in results:
        row_class = ' class="exceeded"' if result['Average_Duration_Minutes'] > 120 else ''
        html_content += f"""
        <tr{row_class}>
            <td>{result['System']}</td>
            <td>{result['Total_Incidents']}</td>
            <td>{result['Average_Duration_Minutes']}</td>
            <td>{result['Average_Duration_Hours']}</td>
            <td>{result['Status']}</td>
        </tr>"""
    
    html_content += """
    </table>
</body>
</html>"""
    
    return html_content

def main():
    excel_file_path = r"C:\Users\XaniyevQX\Desktop\AUTOMATION_NEW\Incident_Report_for_GRC__KRI_ (1).xlsx"
    
    if not os.path.exists(excel_file_path):
        print("Excel file not found!")
        return
    
    # Analyze KRI19
    results = analyze_kri19_from_excel(excel_file_path)
    
    # Generate HTML report
    html_content = generate_kri19_html_report(results)
    
    # Write to file
    with open('KRI19_Average_Resolution_Report.html', 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print("Successfully reported!")

if __name__ == "__main__":
    main()